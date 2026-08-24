"""S4a upload-target, raw-header and non-blocking warning bites-checks."""
from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import sys
import unittest
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1] / ".s4a-test-runtime"
if ROOT.exists():
    shutil.rmtree(ROOT)
ROOT.mkdir(parents=True)
os.environ["SYSTEM_DB_PATH"] = str(ROOT / "system.db")
os.environ["UPLOAD_DIR"] = str(ROOT / "uploads")
os.environ.pop("SYSTEM_DB_BACKUP_PATH", None)

import system_db as s  # noqa: E402
from ai.v2 import service  # noqa: E402
from analysis_runtime.artifacts import AnalysisArtifactRepository  # noqa: E402
from assets import identity  # noqa: E402
from ingest.errors import IngestCorruptionError  # noqa: E402
from ingest.headers import read_raw_headers  # noqa: E402


class S4aUploadTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        s.init_schema()

    def _item(self, kind="dataset"):
        return service.create_item(kind, f"s4a-{self._testMethodName[:35]}")

    def test_next_id_preview_does_not_allocate(self):
        scope = "asset_dataset"
        before = identity.preview_next_id(scope)
        row_before = s.query_one("id_sequences", scope=scope)
        self.assertEqual(identity.preview_next_id(scope), before)
        self.assertEqual(s.query_one("id_sequences", scope=scope), row_before)
        self.assertEqual(identity.allocate(scope), before)

    def test_raw_duplicate_header_blocks_at_step_three_and_names_duplicate(self):
        raw = b"id,id,amount\n1,2,10\n3,4,20\n"
        item = self._item()
        service.save_file(item["item_id"], "data", "duplicate.csv", raw)
        path = Path(s.query("dq_item_files", item_id=item["item_id"], role="data")[0]["path"])
        self.assertEqual(read_raw_headers(path), {"duplicate": ["id", "id", "amount"]})
        self.assertEqual(list(pd.read_csv(io.BytesIO(raw)).columns), ["id", "id.1", "amount"])
        service.profile_item(item["item_id"])
        summary = service.ingest_summary(item["item_id"])
        self.assertEqual(summary["status"], "failed")
        self.assertIn("duplicate column 'id'", summary["fail_reason"])
        self.assertIn("positions 1 and 2", summary["fail_reason"])

    def test_pandas_mangling_counterproof_has_no_visible_duplicate(self):
        raw = b"id,id,amount\n1,2,10\n"
        parsed = list(pd.read_csv(io.BytesIO(raw)).columns)
        self.assertEqual(parsed, ["id", "id.1", "amount"])
        self.assertEqual(parsed.count("id"), 1)
        self.assertIn("id.1", parsed)

    def test_blank_header_blocks_at_step_three(self):
        item = self._item()
        service.save_file(item["item_id"], "data", "blank.csv", b"id, ,amount\n1,2,3\n")
        service.profile_item(item["item_id"])
        summary = service.ingest_summary(item["item_id"])
        self.assertEqual(summary["status"], "failed")
        self.assertIn("blank column header", summary["fail_reason"])

    def test_headerless_numeric_first_row_is_blocked_as_step_two_structure(self):
        item = self._item()
        with self.assertRaises(IngestCorruptionError) as ctx:
            service.save_file(item["item_id"], "data", "headerless.csv", b"1,2\n3,4\n")
        self.assertIn("header row", str(ctx.exception))

    def test_same_header_on_different_workbook_sheets_is_allowed(self):
        from openpyxl import Workbook

        workbook = Workbook()
        first = workbook.active
        first.title = "loans"
        first.append(["id", "amount"])
        first.append(["L1", 10])
        second = workbook.create_sheet("collateral")
        second.append(["id", "value"])
        second.append(["C1", 20])
        payload = io.BytesIO()
        workbook.save(payload)
        item = self._item("database")
        service.save_file(item["item_id"], "data", "source.xlsx", payload.getvalue())
        service.profile_item(item["item_id"])
        self.assertEqual(service.ingest_summary(item["item_id"])["status"], "ready")

    def test_upl11_warnings_are_structured_and_non_blocking(self):
        item = self._item()
        raw = b"all_null,mixed,parse_fail\n,1,2\n,not-a-number,3\n,4,bad\n"
        service.save_file(item["item_id"], "data", "warnings.csv", raw)
        service.save_file(item["item_id"], "dictionary", "types.csv",
                          b"column,declared_type\nparse_fail,number\n")
        service.profile_item(item["item_id"])
        summary = service.ingest_summary(item["item_id"])
        codes = {warning["code"] for warning in summary["warnings"]}
        self.assertTrue({"column_all_null", "column_mixed_type", "column_parse_failure_rate"} <= codes)
        self.assertTrue(all({"column", "code", "message"} <= set(warning) for warning in summary["warnings"]))
        self.assertEqual(summary["status"], "ready")

    def test_source_bundle_applies_parser_controls_and_retains_context(self):
        item = self._item()
        result = service.save_source_bundle(
            item["item_id"], "pipe.txt", b"account|amount\nA1|10\nA2|20\n",
            parsing_options={"delimiter": "pipe"},
            file_context="Quarterly credit-risk population.",
        )
        self.assertEqual(result["summary"][0]["columns"], 2)
        service.profile_item(item["item_id"])
        summary = service.ingest_summary(item["item_id"])
        self.assertEqual(summary["source_parsing_options"], {"delimiter": "pipe"})
        self.assertEqual(summary["file_context"]["text"], "Quarterly credit-risk population.")
        self.assertEqual(
            [row["column_name"] for row in service.get_inventory(item["item_id"])],
            ["account", "amount"],
        )

    def test_workbook_inspection_discovers_data_dictionary_and_editable_context(self):
        from openpyxl import Workbook

        workbook = Workbook()
        data = workbook.active
        data.title = "Portfolio Data"
        data.append(["account", "amount"])
        data.append(["A1", 10])
        dictionary = workbook.create_sheet("Data Dictionary")
        dictionary.append(["column_name", "logical_type", "description"])
        dictionary.append(["amount", "numeric", "Exposure amount"])
        overview = workbook.create_sheet("Overview")
        overview.append(["Purpose", "Quarterly monitoring"])
        overview.append(["Owner", "Credit Risk"])
        payload = io.BytesIO()
        workbook.save(payload)

        inspection = service.inspect_source_upload("portfolio.xlsx", payload.getvalue())
        self.assertEqual(inspection["data_sheet"], "Portfolio Data")
        self.assertEqual(inspection["dictionary_sheet"], "Data Dictionary")
        self.assertEqual(inspection["overview_sheet"], "Overview")
        self.assertIn("Purpose: Quarterly monitoring", inspection["file_context"])

    def test_source_bundle_persists_pre_ingestion_dictionary_header_overrides(self):
        item = self._item()
        service.save_source_bundle(
            item["item_id"], "source.csv", b"amount\n10\n20\n",
            "dictionary.csv",
            b"Field Name,Data Format,Allowed Levels,Missing Codes,Notes\namount,number,0-100,-1,Approved exposure\n",
            dictionary_header_mapping={
                "column_name": "Field Name",
                "logical_type": "Data Format",
                "valid_values": "Allowed Levels",
                "missing_value_codes": "Missing Codes",
                "business_context": "Notes",
            },
        )
        service.profile_item(item["item_id"])
        stored = s.query_one("dq_items", item_id=item["item_id"])
        self.assertTrue(stored["dictionary_mapping_confirmed"])
        self.assertEqual(stored["dictionary_header_mapping_json"]["valid_values"], "Allowed Levels")
        version = s.query_one(
            "dq_asset_dictionaries", dictionary_version_id=stored["dictionary_version_id"]
        )
        self.assertEqual(version["parsed_json"]["*"]["amount"]["valid_values"], "0-100")

    def test_quarter_profile_uses_full_column_for_calendar_bounds(self):
        item = self._item()
        service.save_source_bundle(
            item["item_id"], "quarters.csv",
            b"reporting_period,value\n2006Q2,1\n2007Q4,2\n2006Q1,3\n2007Q1,4\n",
            "dictionary.csv",
            b"Field Name,Data Format,Requirement\nreporting_period,period,period\nvalue,number,target\n",
            dictionary_header_mapping={"column_name": "Field Name", "logical_type": "Data Format",
                                       "role": "Requirement"},
        )
        service.profile_item(item["item_id"])

        period = next(row for row in service.get_inventory(item["item_id"])
                      if row["column_name"] == "reporting_period")
        self.assertEqual(period["dictionary_role"], "period")
        self.assertEqual(period["role"], "Period")
        target = next(row for row in service.get_inventory(item["item_id"])
                      if row["column_name"] == "value")
        self.assertEqual(target["dictionary_role"], "target")
        self.assertEqual(target["role"], "Target")
        self.assertEqual(period["profile_json"]["min"], "2006Q1")
        self.assertEqual(period["profile_json"]["max"], "2007Q4")
        self.assertEqual(period["profile_json"]["period_bounds"], {
            "start_date": "2006-01-01",
            "end_date": "2007-12-31",
            "format": "calendar_quarter",
        })

    def test_ready_snapshot_and_dictionary_survive_application_restart(self):
        item = self._item()
        service.save_source_bundle(
            item["item_id"], "restart.csv", b"default_flag,amount\n0,10\n1,20\n",
            "dictionary.csv",
            b"column,definition,declared_type,role\ndefault_flag,Outcome,binary,target\namount,Exposure,number,feature\n",
        )
        service.profile_item(item["item_id"])
        service.process_snapshot(item["item_id"], snapshot_label="restart-proof",
                                 target_variable="default_flag", use_case="IFRS 9", product="CRE")

        probe = """
import json
import system_db as s
item_id = {item_id!r}
snapshot = s.query_one('dq_items', item_id=item_id)
inventory = s.query('variable_inventory', item_id=item_id)
dictionary = s.query_one('dq_asset_dictionaries', dictionary_version_id=snapshot['dictionary_version_id'])
print(json.dumps({{
    'status': snapshot['ingest_status'],
    'label': snapshot['snapshot_label'],
    'target': snapshot['target_variable'],
    'dictionary_asset_id': dictionary['asset_id'],
    'columns': sorted(row['column_name'] for row in inventory),
}}))
""".format(item_id=item["item_id"])
        restart_env = {**os.environ, "SYSTEM_DB_PATH": str(s.SYS_DB_PATH),
                       "UPLOAD_DIR": str(ROOT / "uploads")}
        restart_env.pop("SYSTEM_DB_BACKUP_PATH", None)
        restarted = subprocess.run(
            [sys.executable, "-c", probe], cwd=Path(__file__).resolve().parents[1],
            env=restart_env,
            check=True, capture_output=True, text=True,
        )
        persisted = json.loads(restarted.stdout.strip().splitlines()[-1])
        self.assertEqual(persisted, {
            "status": "ready",
            "label": "restart-proof",
            "target": "default_flag",
            "dictionary_asset_id": item["asset_id"],
            "columns": ["amount", "default_flag"],
        })

    def test_confirmed_type_map_is_on_snapshot_and_fresh_reference_schema(self):
        item = self._item()
        service.save_file(item["item_id"], "data", "types.csv", b"amount,label\n1,a\n2,b\n")
        service.profile_item(item["item_id"])
        rows = service.get_inventory(item["item_id"])
        rows[0]["classification"] = "text"
        service.put_inventory(item["item_id"], rows)
        snapshot = s.query_one("dq_items", item_id=item["item_id"])
        asset = s.query_one("dq_assets", asset_id=item["asset_id"])
        version = s.query_one("dq_asset_versions", asset_id=asset["asset_id"], version_no=1)
        snapshot_map = snapshot["column_type_map_json"] if isinstance(snapshot["column_type_map_json"], dict) else json.loads(snapshot["column_type_map_json"])
        reference = version["reference_schema_json"] if isinstance(version["reference_schema_json"], dict) else json.loads(version["reference_schema_json"])
        self.assertEqual(snapshot_map, reference)
        self.assertEqual(snapshot_map["tables"]["types"]["types"][rows[0]["column_name"]], "text")
        self.assertEqual(version["created_from_snapshot_id"], item["item_id"])

    def test_process_persists_reviewed_roles_in_saved_schema_artifacts(self):
        item = self._item()
        service.save_file(
            item["item_id"], "data", "roles.csv",
            b"DPD,non_accrual_indicator,exposure\n0,0,10\n30,0,20\n",
        )
        service.profile_item(item["item_id"])
        rows = service.get_inventory(item["item_id"])
        for row in rows:
            if row["column_name"] in {"DPD", "non_accrual_indicator"}:
                row["role"] = "Ignore"

        service.process_snapshot(
            item["item_id"], snapshot_label="reviewed-schema",
            target_variable=None, use_case="IFRS 9", product="CRE",
            inventory_rows=rows,
        )

        inventory = {row["column_name"]: row for row in service.get_inventory(item["item_id"])}
        self.assertEqual(inventory["DPD"]["role"], "Ignore")
        self.assertEqual(inventory["non_accrual_indicator"]["role"], "Ignore")
        artifacts = AnalysisArtifactRepository().list(
            snapshot_id=item["item_id"], artifact_type="column_profile", status="active"
        )
        saved_roles = {artifact.feature: artifact.summary["fields"]["role"] for artifact in artifacts}
        self.assertEqual(saved_roles["DPD"], "Ignore")
        self.assertEqual(saved_roles["non_accrual_indicator"], "Ignore")


if __name__ == "__main__":
    unittest.main()
