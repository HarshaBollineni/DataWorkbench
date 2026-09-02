"""Advisory FNC checker.

This is intentionally informational: it reports defects but never asserts or
raises on them, so FNC-04 cannot gate ci-local.ps1 or any release build.
"""
from __future__ import annotations

import ast
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


WORKTREE = Path(__file__).resolve().parents[2]


def _find_workbook() -> Path:
    """Locate the advisory workbook without assuming checkout nesting depth."""
    roots = (WORKTREE, *WORKTREE.parents)
    return next(
        (root / "functional_tools.xlsx" for root in roots
         if (root / "functional_tools.xlsx").is_file()),
        WORKTREE.parent / "functional_tools.xlsx",
    )


WORKBOOK = _find_workbook()
HEADERS = ["Sl. No.", "Important functions", "Function purpose", "Inputs", "Processing", "Outputs"]
SHEET_SOURCES = {
    "service.py": ["ai/v2/service.py", "assets/service.py"],
    "diagnostics-core": [
        "dq_diagnostics/register.py", "dq_diagnostics/runner.py", "dq_diagnostics/guards.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/manifest.py",
        "dq_diagnostics/readiness.py", "dq_diagnostics/result.py",
        "dq_diagnostics/thresholds.py", "dq_diagnostics/delivery.py",
        "dq_diagnostics/profiling_preconditions.py",
    ],
    "diagnostics-engines": [
        "dq_diagnostics/engines/base.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/binder.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/engine.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/primitives.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/result.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/roles.py",
        "domains/test_lab/diagnostics/t2_d04_cross_field_business_rule/rules.py",
    ],
    "dq_tests": ["dq_tests/contracts.py", "dq_tests/global_rules.py", "dq_tests/param_specs.py", "dq_tests/registry.py"],
    "gx-scoring": ["gx/context.py", "gx/runner.py", "gx/suite_builder.py", "gx/metrics.py", "scoring/health.py", "scoring/criticality.py", "scoring/categories.py"],
    "kb": ["kb.py", "kb_convert.py"],
    "rca": [
        "domains/rca/service.py", "domains/rca/effective_challenge.py",
        "domains/rca/analysis_helpers.py",
    ],
    "issues": ["ai/v2/issues.py"],
    "taxonomy": ["taxonomy.py", "seeds/taxonomy_seed.py"],
    "admin": ["routers/admin.py", "system_db.py", "tenancy.py"],
    "auth": ["routers/auth.py"],
    "assets": ["assets/identity.py", "assets/service.py", "assets/reads.py", "assets/schema_check.py", "assets/diff.py", "assets/refresh.py", "assets/staleness.py", "assets/history.py"],
    "analytics": ["analytics/events.py", "analytics/measures.py", "analytics/ids.py", "analytics/catalogue.py"],
}


def _functions(path: Path) -> dict[str, list[int]]:
    try:
        tree = ast.parse(path.read_text(encoding="utf-8"))
    except Exception as exc:  # advisory only
        return {f"<parse error: {exc}>": []}
    return {
        node.name: [node.lineno, getattr(node, "end_lineno", node.lineno)]
        for node in ast.walk(tree)
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
    }


def findings() -> list[str]:
    out: list[str] = []
    if not WORKBOOK.exists():
        return [f"workbook missing: {WORKBOOK}"]
    wb = load_workbook(WORKBOOK, read_only=False, data_only=False)
    if set(wb.sheetnames) != set(SHEET_SOURCES):
        out.append(f"sheet set differs: found {wb.sheetnames!r}")
    backend = WORKTREE / "backend"
    for name, sources in SHEET_SOURCES.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if ws.max_column != 6:
            out.append(f"{name}: expected six columns, found {ws.max_column}")
        header = [ws.cell(4, c).value for c in range(1, 7)]
        if header != HEADERS:
            out.append(f"{name}: header differs: {header!r}")
        if not ws["A1"].value or "\n" not in str(ws["A2"].value or ""):
            out.append(f"{name}: two-line header block is incomplete")
        source_functions = {}
        for rel in sources:
            path = backend / rel
            if path.exists():
                source_functions[rel] = _functions(path)
        for r in range(5, ws.max_row + 1):
            cell = str(ws.cell(r, 2).value or "")
            match = re.match(r"(?:[\w]+\.)*([A-Za-z_]\w*)\s*\(", cell)
            if not match:
                continue  # documentation marker, such as _MIGRATIONS or retired stage shims
            fn = match.group(1)
            if fn == "_MIGRATIONS":
                continue  # documented data structure, not a callable
            line_match = re.search(r"lines\s+(\d+)-(\d+)", cell)
            candidates = [(rel, funcs[fn]) for rel, funcs in source_functions.items() if fn in funcs]
            if not candidates:
                out.append(f"{name} row {r}: {fn} not found in claimed sources")
                continue
            if line_match:
                recorded = int(line_match.group(1))
                if not any(abs(recorded - loc[0]) <= 10 for _, loc in candidates):
                    out.append(f"{name} row {r}: {fn} line {recorded} is outside ±10 of live source")
    out.append("helper/tool layer intentionally absent: PLT-08 docstring contracts already cover it; no duplicate sheet was added")
    return out


def test_functional_map_advisory() -> None:
    """Run informational findings while remaining non-blocking by design."""
    for item in findings():
        print(f"[FNC advisory] {item}")


if __name__ == "__main__":
    for item in findings():
        print(f"[FNC advisory] {item}")
    sys.exit(0)
