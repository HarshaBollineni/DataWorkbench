"""RCA Stage 2 — Knowledge Base behavioral tests.

Standalone-runnable (``python -m unittest tests.test_kb``); see
test_taxonomy.py's module docstring for the SYSTEM_DB_PATH-at-import-time
constraint this file follows the same way. KB_STORAGE_DIR is likewise
sandboxed to a temp directory so tests never write into backend/kb_storage.
"""
from __future__ import annotations

import os
import shutil
import sqlite3
import tempfile
import unittest
import uuid
from pathlib import Path
from unittest import mock

_TMP_DB = Path(tempfile.gettempdir()) / "archimedes-test-kb.db"
_TMP_KB_STORAGE = Path(tempfile.gettempdir()) / "archimedes-test-kb-storage"
if _TMP_DB.exists():
    _TMP_DB.unlink()
if _TMP_KB_STORAGE.exists():
    shutil.rmtree(_TMP_KB_STORAGE)
os.environ["SYSTEM_DB_PATH"] = str(_TMP_DB)
os.environ["KB_STORAGE_DIR"] = str(_TMP_KB_STORAGE)
os.environ.pop("SYSTEM_DB_BACKUP_PATH", None)

import system_db as s  # noqa: E402
import kb  # noqa: E402
from kb_convert import ConversionError, convert_to_markdown  # noqa: E402
from seeds import backfill_kb_roles  # noqa: E402
from seeds.taxonomy_seed import BOOTSTRAP_TENANT, seed_platform, seed_taxonomy  # noqa: E402

TENANT = BOOTSTRAP_TENANT
EDITOR_ROLES = ["kb_editor", "kb_reviewer"]


def _make_pdf_bytes(text: str | None) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    if text:
        pdf.set_font("Helvetica", "", 12)
        pdf.multi_cell(0, 8, text)
    out = pdf.output()
    return bytes(out)


def _make_docx_bytes(heading: str, body: str) -> bytes:
    import io
    import docx
    document = docx.Document()
    document.add_heading(heading, level=2)
    document.add_paragraph(body)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


class ConversionTests(unittest.TestCase):
    def test_txt_passthrough(self):
        result = convert_to_markdown("notes.txt", "text/plain", b"## Rule\nBody text.")
        self.assertIn("Body text.", result["markdown"])
        self.assertEqual(result["converter_name"], "passthrough-txt")

    def test_md_passthrough(self):
        result = convert_to_markdown("notes.md", "text/markdown", b"## Rule\nBody text.")
        self.assertEqual(result["converter_name"], "passthrough-md")

    def test_docx_conversion(self):
        content = _make_docx_bytes("A heading", "A body paragraph.")
        result = convert_to_markdown("doc.docx", "", content)
        self.assertIn("## A heading", result["markdown"])
        self.assertIn("A body paragraph.", result["markdown"])
        self.assertEqual(result["converter_name"], "python-docx")

    def test_pdf_with_text_converts(self):
        content = _make_pdf_bytes("This PDF has a real text layer with enough characters.")
        result = convert_to_markdown("doc.pdf", "", content)
        self.assertIn("real text layer", result["markdown"])
        self.assertEqual(result["converter_name"], "pdfplumber")

    def test_pdf_without_text_is_rejected(self):
        content = _make_pdf_bytes(None)  # blank page — no text layer
        with self.assertRaises(ConversionError):
            convert_to_markdown("scan.pdf", "", content)

    def test_unsupported_extension_rejected(self):
        with self.assertRaises(ConversionError):
            convert_to_markdown("archive.zip", "", b"PK\x03\x04")

    def test_empty_file_rejected(self):
        with self.assertRaises(ConversionError):
            convert_to_markdown("empty.txt", "text/plain", b"")


class KbLifecycleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        s.init_schema()
        seed_platform()
        seed_taxonomy()

    def _upload_and_submit(self, filename="rules.md", body="## Heading one\nBody one.\n\n## Heading two\nBody two.\n"):
        upload = kb.upload_document(TENANT, filename, "text/markdown", body.encode("utf-8"),
                                    "domain_fact", "tester")
        preview = kb.submit_for_review(TENANT, upload["version"]["version_id"], "tester", "domain_fact")
        return upload, preview

    def _draft_proposal(self, subject: str, *, based_on_rule_id: str | None = None,
                        source: str = "case_closure", conn=None) -> dict:
        metadata = {
            "proposal_kind": "t2_d11_expected_direction",
            "proposal_subject_key": subject,
            "proposal_decision_hash": f"decision-{subject}",
            "based_on_rule_id": based_on_rule_id,
        }
        rule = kb.draft_rule_from_case_closure(
            TENANT, "domain_fact", f"Directionality proposal {subject}",
            ["portfolio"], f"run-{subject}", "tester",
            proposal_metadata=metadata, conn=conn,
        )
        if source != "case_closure":
            version = s.query_one("kb_document_versions", version_id=rule["version_id"])
            report = version["conversion_report_json"]
            s.update("kb_document_versions", {"version_id": rule["version_id"]}, {
                "conversion_report_json": {**report, "source": source},
            })
        return rule

    def test_upload_preserves_original_bytes_and_hash(self):
        content = b"## Heading\nBody."
        upload = kb.upload_document(TENANT, "a.md", "text/markdown", content, "domain_fact", "tester")
        version = upload["version"]
        self.assertEqual(kb.read_original(version["original_bytes_ref"]), content)
        import hashlib
        self.assertEqual(version["original_sha256"], hashlib.sha256(content).hexdigest())

    def test_upload_storage_is_content_addressed_dedup(self):
        content = b"## Heading\nIdentical body for dedup test."
        u1 = kb.upload_document(TENANT, "a.md", "text/markdown", content, "domain_fact", "tester")
        u2 = kb.upload_document(TENANT, "b.md", "text/markdown", content, "domain_fact", "tester")
        self.assertEqual(u1["version"]["original_bytes_ref"], u2["version"]["original_bytes_ref"])

    def test_submit_for_review_creates_one_draft_rule_per_section(self):
        _, preview = self._upload_and_submit()
        self.assertEqual(len(preview["sections"]), 2)
        self.assertEqual(len(preview["rules"]), 2)
        self.assertTrue(all(r["lifecycle_state"] == "draft" for r in preview["rules"]))
        self.assertTrue(all(r["trust_level"] == "inferred" for r in preview["rules"]))

    def test_publish_requires_kb_reviewer_role(self):
        _, preview = self._upload_and_submit()
        rule_id = preview["rules"][0]["rule_id"]
        with self.assertRaises(kb.ForbiddenError):
            kb.publish_rule(TENANT, rule_id, "tester", roles=["kb_editor"], category="domain_fact")

    def test_publish_succeeds_with_role_and_sets_effective_fields(self):
        _, preview = self._upload_and_submit()
        rule_id = preview["rules"][0]["rule_id"]
        published = kb.publish_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, category="domain_fact")
        self.assertEqual(published["lifecycle_state"], "published")
        self.assertEqual(published["trust_level"], "human_confirmed")
        self.assertIsNotNone(published["effective_date"])
        self.assertEqual(published["shelf_life_months"], kb.DEFAULT_SHELF_LIFE_MONTHS)

    def test_draft_only_rules_never_appear_in_retrieval(self):
        _, preview = self._upload_and_submit()
        result = kb.list_eligible_rules(TENANT, "intake")
        rule_ids = {r["rule_id"] for r in result["rules"]}
        self.assertFalse(rule_ids & {r["rule_id"] for r in preview["rules"]})

    def test_diagnostic_proposal_stays_in_learning_candidate_queue(self):
        subject = f"subject-{uuid.uuid4().hex}"
        rule = self._draft_proposal(subject, source="diagnostic_proposal")

        self.assertIn(rule["version_id"], kb.learning_candidate_version_ids(TENANT))
        self.assertIn(rule["rule_id"], {
            row["candidate_id"] for row in kb.list_learning_candidates(TENANT)
        })

    def test_proposal_hierarchy_rolls_back_when_unique_subject_insert_loses(self):
        subject = f"subject-{uuid.uuid4().hex}"
        self._draft_proposal(subject)
        before = {
            table: len(s.query(table))
            for table in ("kb_documents", "kb_document_versions", "kb_sections", "kb_rules")
        }

        with self.assertRaises(sqlite3.IntegrityError):
            self._draft_proposal(subject)

        after = {
            table: len(s.query(table))
            for table in ("kb_documents", "kb_document_versions", "kb_sections", "kb_rules")
        }
        self.assertEqual(after, before)

        with s.get_conn() as conn:
            with self.assertRaises(sqlite3.IntegrityError):
                self._draft_proposal(subject, conn=conn)
            # The caller catches the unique-index loser and safely continues
            # its wider transaction; the proposal savepoint must have removed
            # all three hierarchy parents created before the rule insert.
            s.insert("transaction_log", {
                "ts": s.now_ist(), "actor": "tester",
                "event": "proposal_savepoint_continued", "payload": {},
            }, conn=conn)
            conn.commit()
        self.assertEqual({
            table: len(s.query(table))
            for table in ("kb_documents", "kb_document_versions", "kb_sections", "kb_rules")
        }, before)

        rollback_subject = f"subject-{uuid.uuid4().hex}"
        with s.get_conn() as conn:
            self._draft_proposal(rollback_subject, conn=conn)
            conn.rollback()
        self.assertEqual({
            table: len(s.query(table))
            for table in ("kb_documents", "kb_document_versions", "kb_sections", "kb_rules")
        }, before)

    def test_linked_revision_requires_same_published_proposal_subject(self):
        subject = f"subject-{uuid.uuid4().hex}"
        predecessor = self._draft_proposal(subject)
        kb.publish_rule(
            TENANT, predecessor["rule_id"], "reviewer", EDITOR_ROLES, "domain_fact",
        )
        wrong_subject = self._draft_proposal(
            f"other-{uuid.uuid4().hex}", based_on_rule_id=predecessor["rule_id"],
        )

        with self.assertRaisesRegex(kb.KbError, "same proposal subject"):
            kb.publish_rule(
                TENANT, wrong_subject["rule_id"], "reviewer", EDITOR_ROLES,
                "domain_fact",
            )

        self.assertEqual(
            s.query_one("kb_rules", rule_id=predecessor["rule_id"])["lifecycle_state"],
            "published",
        )
        self.assertEqual(
            s.query_one("kb_rules", rule_id=wrong_subject["rule_id"])["lifecycle_state"],
            "draft",
        )

    def test_linked_revision_publish_and_supersession_roll_back_together(self):
        subject = f"subject-{uuid.uuid4().hex}"
        predecessor = self._draft_proposal(subject)
        kb.publish_rule(
            TENANT, predecessor["rule_id"], "reviewer", EDITOR_ROLES, "domain_fact",
        )
        revision = self._draft_proposal(
            subject, based_on_rule_id=predecessor["rule_id"],
        )
        real_insert = s.insert

        def fail_supersession_audit(table, row, *, conn=None):
            if table == "transaction_log" and row.get("event") == "knowledge_superseded":
                raise RuntimeError("simulated audit write failure")
            return real_insert(table, row, conn=conn)

        with mock.patch.object(kb.s, "insert", side_effect=fail_supersession_audit):
            with self.assertRaisesRegex(RuntimeError, "audit write failure"):
                kb.publish_rule(
                    TENANT, revision["rule_id"], "reviewer", EDITOR_ROLES,
                    "domain_fact",
                )

        self.assertEqual(
            s.query_one("kb_rules", rule_id=predecessor["rule_id"])["lifecycle_state"],
            "published",
        )
        self.assertEqual(
            s.query_one("kb_rules", rule_id=revision["rule_id"])["lifecycle_state"],
            "draft",
        )
        self.assertFalse(any(
            row["event"] == "knowledge_publish"
            and (row.get("payload") or {}).get("rule_id") == revision["rule_id"]
            for row in s.query("transaction_log")
        ))

        kb.publish_rule(
            TENANT, revision["rule_id"], "reviewer", EDITOR_ROLES, "domain_fact",
        )
        self.assertEqual(
            s.query_one("kb_rules", rule_id=predecessor["rule_id"])["lifecycle_state"],
            "superseded",
        )
        self.assertEqual(
            s.query_one("kb_rules", rule_id=revision["rule_id"])["lifecycle_state"],
            "published",
        )

    def test_new_candidate_can_publish_without_reopening_archived_predecessor(self):
        subject = f"subject-{uuid.uuid4().hex}"
        predecessor = self._draft_proposal(subject)
        kb.archive_rule(
            TENANT, predecessor["rule_id"], "reviewer", EDITOR_ROLES,
            "Reviewer rejected the earlier wording.",
        )
        replacement = self._draft_proposal(
            subject, based_on_rule_id=predecessor["rule_id"],
        )

        published = kb.publish_rule(
            TENANT, replacement["rule_id"], "reviewer", EDITOR_ROLES,
            "domain_fact",
        )

        self.assertEqual(published["lifecycle_state"], "published")
        self.assertEqual(
            s.query_one("kb_rules", rule_id=predecessor["rule_id"])["lifecycle_state"],
            "archived",
        )
        self.assertIsNone(
            s.query_one("kb_rules", rule_id=predecessor["rule_id"])[
                "superseded_by_rule_id"
            ]
        )

    def test_archive_requires_reason_and_role(self):
        _, preview = self._upload_and_submit()
        rule_id = preview["rules"][0]["rule_id"]
        kb.publish_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, category="domain_fact")
        with self.assertRaises(kb.ForbiddenError):
            kb.archive_rule(TENANT, rule_id, "tester", roles=[], reason="stale")
        with self.assertRaises(ValueError):
            kb.archive_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, reason="")
        archived = kb.archive_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, reason="superseded")
        self.assertEqual(archived["lifecycle_state"], "archived")

    def test_shelf_life_expiry_flips_trust_level(self):
        _, preview = self._upload_and_submit()
        rule_id = preview["rules"][0]["rule_id"]
        kb.publish_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, category="domain_fact",
                        shelf_life_months=1)
        s.update("kb_rules", {"rule_id": rule_id}, {"last_confirmed_date": _months_ago(3)})
        flipped = kb.check_shelf_life(TENANT)
        self.assertGreaterEqual(flipped, 1)
        self.assertEqual(s.query_one("kb_rules", rule_id=rule_id)["trust_level"], "inferred")

    def test_schema_change_invalidation_flips_trust_level(self):
        # table_metadata has no natural unique key besides its autoincrement
        # id, so s.upsert() here would INSERT a second row rather than
        # replace — use s.update() to mutate the existing row in place, the
        # way a real re-profile of a table would.
        s.insert("table_metadata", {"logical_db": "test_db", "table": "kb_test_table",
                                    "columns": ["a", "b"], "datatypes": ["int", "str"]})
        _, preview = self._upload_and_submit()
        rule_id = preview["rules"][0]["rule_id"]
        kb.publish_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, category="domain_fact",
                        related_tables=["kb_test_table"])
        self.assertEqual(kb.check_schema_invalidation(TENANT), 0)  # no change yet
        s.update("table_metadata", {"logical_db": "test_db", "table": "kb_test_table"},
                {"columns": ["a", "b", "c"], "datatypes": ["int", "str", "float"]})
        flipped = kb.check_schema_invalidation(TENANT)
        self.assertGreaterEqual(flipped, 1)
        self.assertEqual(s.query_one("kb_rules", rule_id=rule_id)["trust_level"], "inferred")

    def test_mark_under_suspicion(self):
        _, preview = self._upload_and_submit()
        rule_id = preview["rules"][0]["rule_id"]
        kb.publish_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, category="domain_fact")
        flagged = kb.mark_under_suspicion(TENANT, rule_id, "case XYZ ended Unresolved")
        self.assertEqual(flagged["lifecycle_state"], "under_suspicion")
        self.assertIn("Unresolved", flagged["under_suspicion_reason"])

    def test_retrieval_respects_agent_category_matrix_and_quarantine(self):
        _, preview = self._upload_and_submit(
            filename="hist.md", body="## Case history rule\nAn accepted expected change.\n")
        rule_id = preview["rules"][0]["rule_id"]
        kb.publish_rule(TENANT, rule_id, "tester", roles=EDITOR_ROLES, category="case_history")

        planner_result = kb.list_eligible_rules(TENANT, "planner")
        self.assertNotIn(rule_id, {r["rule_id"] for r in planner_result["rules"]})  # quarantined

        pass2_result = kb.list_eligible_rules(TENANT, "coverage_challenge_pass2")
        self.assertIn(rule_id, {r["rule_id"] for r in pass2_result["rules"]})  # permitted to nominate

        manifest = s.query_one("kb_retrieval_manifests", manifest_id=pass2_result["manifest_id"])
        self.assertIn(rule_id, manifest["rule_ids_json"])

    def test_unknown_agent_gets_nothing(self):
        result = kb.list_eligible_rules(TENANT, "not_a_real_agent")
        self.assertEqual(result["rules"], [])


def _months_ago(n: float) -> str:
    from datetime import datetime, timedelta
    return (datetime.fromisoformat(s.now_ist()) - timedelta(days=n * 30.436875)).isoformat()


class SyntheticKbTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        s.init_schema()
        seed_platform()
        seed_taxonomy()

    def test_never_loaded_on_boot_then_explicit_load_produces_expected_states(self):
        """Single ordered test (not two independent test_ methods) so the
        "never loaded on a normal boot" assertion is structurally guaranteed
        to run before anything calls load_synthetic_kb() in this process —
        unittest has no cross-method ordering guarantee within a class, and
        splitting this into two methods let a prior loader test's leftover
        rows produce a false failure here."""
        # Mirrors main.py's boot sequence (init_schema + seed_platform_and_
        # taxonomy, already done in setUpClass): zero synthetic kb_rules
        # without an explicit load_synthetic_kb() call.
        self.assertEqual([r for r in s.query("kb_rules") if r.get("is_synthetic")], [])

        from seeds.synthetic_kb_seed import load_synthetic_kb
        result = load_synthetic_kb(TENANT, actor="test-loader")
        self.assertEqual(result["documents"], 5)
        self.assertGreaterEqual(result["rules_published"], 8)
        self.assertEqual(result["rules_left_draft"], 1)

        synthetic_rules = [r for r in s.query("kb_rules", tenant_id=TENANT) if r.get("is_synthetic")]
        self.assertTrue(all(r["is_synthetic"] for r in synthetic_rules))

        outcome_leakage = next(r for r in synthetic_rules if "default indicator" in r["rule_text"])
        # Backdated 18mo against a 12mo shelf life at load time; the loader
        # itself does not run the expiry sweep, so this only becomes
        # 'inferred' once check_shelf_life() actually runs — confirming the
        # manifest's last_confirmed_months_ago directive really does produce
        # a shelf-life-exceeded rule, not just a backdated timestamp.
        self.assertEqual(outcome_leakage["trust_level"], "human_confirmed")
        kb.check_shelf_life(TENANT)
        outcome_leakage_after = s.query_one("kb_rules", rule_id=outcome_leakage["rule_id"])
        self.assertEqual(outcome_leakage_after["trust_level"], "inferred")

        # Every document also gets an auto-generated "Overview" draft rule
        # for any content preceding its first "## " heading (here, just the
        # "# Title" line) — so lifecycle_state == "draft" alone is ambiguous;
        # match the specific manifest-declared draft to check it stayed one.
        draft_rule = next(r for r in synthetic_rules if "intentionally left as a draft" in r["rule_text"])
        self.assertEqual(draft_rule["lifecycle_state"], "draft")

        suspicious = [r for r in synthetic_rules if r["lifecycle_state"] == "under_suspicion"]
        self.assertEqual(len(suspicious), 1)


class BackfillKbRolesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        s.init_schema()

    def test_backfill_unions_configured_kb_roles_without_clobbering_others(self):
        s.upsert("users", {"username": "kb_backfill_test", "password": "x",
                           "authz_roles": ["user", "some_custom_role"]})
        import seeds
        original_cfg = seeds._users_config
        seeds._users_config = lambda: {"users": [
            {"username": "kb_backfill_test", "authz_roles": ["user", "admin", "kb_editor", "kb_reviewer"]},
        ]}
        try:
            n = backfill_kb_roles()
            self.assertEqual(n, 1)
            row = s.query_one("users", username="kb_backfill_test")
            self.assertEqual(set(row["authz_roles"]),
                             {"user", "some_custom_role", "kb_editor", "kb_reviewer"})
            self.assertEqual(backfill_kb_roles(), 0)  # idempotent
        finally:
            seeds._users_config = original_cfg


if __name__ == "__main__":
    unittest.main()
