"""Raw tabular-header reading for UPL-12.

Pandas is deliberately not used here: ``read_csv`` and ``read_excel``
silently mangle duplicate and blank headers before callers can inspect them.
The returned values preserve the literal header text (apart from converting
Excel cell values to strings); validation strips only for blank/duplicate
comparison and remains case-sensitive, so ``Id`` and ``id`` are different.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path


def _csv_dialect(text: str) -> csv.Dialect:
    try:
        return csv.Sniffer().sniff(text, delimiters=",\t;|")
    except csv.Error:
        return csv.excel


def _csv_headers(path: Path, delimiter: str | None = None) -> list[str]:
    raw = path.read_bytes()
    if not raw:
        return []
    text = raw.decode("utf-8-sig", errors="replace")
    first_line = text.splitlines()[0] if text.splitlines() else ""
    if not first_line:
        return []
    reader = csv.reader([first_line], delimiter=delimiter) if delimiter else csv.reader([first_line], _csv_dialect(first_line))
    return ["" if value is None else str(value) for value in next(reader)]


def _xlsx_headers(path: Path, sheet_name: str | int | None = None) -> dict[str, list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {
            str(sheet.title): ["" if value is None else str(value) for value in next(
                sheet.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )]
            for sheet in workbook.worksheets
        }
        if sheet_name is None:
            return result
        if isinstance(sheet_name, int):
            names = list(result)
            return {names[sheet_name]: result[names[sheet_name]]} if 0 <= sheet_name < len(names) else {}
        return {str(sheet_name): result[str(sheet_name)]} if str(sheet_name) in result else {}
    finally:
        workbook.close()


def read_raw_headers(path: str | Path, *, delimiter: str | None = None,
                     sheet_name: str | int | None = None) -> dict[str, list[str]]:
    """Return literal first-row headers per table/sheet without pandas mangling."""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        selected_delimiter = delimiter or ("\t" if suffix == ".tsv" else None)
        return {source.stem: _csv_headers(source, selected_delimiter)}
    if suffix == ".xlsx":
        return _xlsx_headers(source, sheet_name)
    if suffix == ".xls":
        # The S4a picker is intentionally .csv/.xlsx only. Let the existing
        # pandas reader own legacy .xls compatibility; raw inspection cannot
        # be made reliable here without reintroducing a separate engine.
        return {}
    raise ValueError(f"Unsupported data file type for raw headers: {suffix}")


def validate_raw_headers(headers_by_table: dict[str, list[str]]) -> None:
    """Raise a structural ingestion error for blank/duplicate table headers.

    Comparison is per table, never across workbook sheets. It is
    case-sensitive by design: ``Id`` and ``id`` are distinct columns.
    """
    from ingest.errors import IngestCorruptionError

    if not headers_by_table or any(not headers for headers in headers_by_table.values()):
        raise IngestCorruptionError("The data file has no header row.")
    for table, headers in headers_by_table.items():
        seen: dict[str, int] = {}
        for position, raw in enumerate(headers, start=1):
            name = str(raw).strip()
            if not name:
                raise IngestCorruptionError(
                    f"Table '{table}' has a blank column header at position {position}."
                )
            if name in seen:
                raise IngestCorruptionError(
                    f"Table '{table}' declares duplicate column '{name}' at positions "
                    f"{seen[name]} and {position}; duplicate headers are not allowed."
                )
            seen[name] = position
