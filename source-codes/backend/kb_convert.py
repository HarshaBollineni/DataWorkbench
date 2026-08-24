"""RCA Stage 2 — Knowledge document conversion (contracts.md §6.1).

TXT/Markdown normalize directly. DOCX/PDF convert to Markdown internally via
libraries already present in backend/requirements.txt (python-docx,
pdfplumber — the same ones ai/dict_ingest.py already uses for data-dictionary
extraction). Scanned/image-only PDFs are rejected, never silently accepted as
error-free text (contracts.md §6.1: "never treat OCR output as error-free
domain truth" — here we go further and refuse OCR entirely, per the Stage 0
decision that OCR is out of scope).

Phase 5 (KB-04, docs/0.4.0/04-kb-contract.md §1.3) — table-aware conversion.
Single-provenance-artifact principle: the stored ``converted_markdown`` IS
what downstream parsing (kb.py) sees, so table structure must survive the
PDF/DOCX -> Markdown step as real GFM pipe tables, not get flattened into
prose. This module owns detection+rendering only; interpreting a rendered
table's rows as rule records is kb.py's job (KB-01: no domain content here).
"""
from __future__ import annotations

import hashlib
import io
import re

CONVERTER_VERSION = "1"
SUPPORTED_MEDIA_TYPES = {
    "text/plain": "txt",
    "text/markdown": "md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/pdf": "pdf",
}
MAX_BYTES = 20 * 1024 * 1024  # 20 MB (Stage 0 decision §9.6)

# A PDF whose extracted text density falls below this is treated as
# scanned/image-only (no text layer) and rejected rather than silently
# producing an empty or near-empty rule set.
_PDF_MIN_CHARS_PER_PAGE = 20


class ConversionError(Exception):
    """Unsupported, corrupt, or scanned/image-only source — rejected safely."""


def sniff_media_type(filename: str, declared_media_type: str | None) -> str:
    ext = (filename.rsplit(".", 1)[-1] or "").lower() if "." in filename else ""
    by_ext = {"txt": "text/plain", "md": "text/markdown",
             "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             "pdf": "application/pdf"}
    if ext in by_ext:
        return by_ext[ext]
    if declared_media_type in SUPPORTED_MEDIA_TYPES:
        return declared_media_type
    raise ConversionError(f"Unsupported file type (extension {ext!r}). "
                          f"Supported: .txt, .md, .docx, .pdf")


# --- Table rendering (shared by PDF + DOCX) — generic, no document-specific
# knowledge. A "table" here is any grid of cells a parser (pdfplumber's
# find_tables(), python-docx's document.tables) already identified as one;
# this only cleans cell text and renders GFM syntax so it survives as
# structure in the stored Markdown. -----------------------------------------

def _clean_cell(text: str | None) -> str:
    """Collapse a PDF/DOCX table cell's embedded line-wraps into one line.

    A hyphen immediately followed by a line break is treated as a broken
    token continuing on the next line (e.g. an ID cell wrapped mid-word as
    "IFRS9-\\n01") and rejoined WITHOUT inserting a space -> "IFRS9-01".
    Any other line break (e.g. a long "Roles (semantic)" cell wrapping
    naturally) becomes a single space. Generic string cleanup — no column-
    specific logic.
    """
    if text is None:
        return ""
    text = str(text).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"-\s*\n\s*", "-", text)
    text = text.replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def _escape_pipe(cell: str) -> str:
    return cell.replace("|", "\\|")


def _rows_to_gfm(rows: list[list] | None) -> str:
    """Render a raw cell-matrix (first row = header) as a GFM pipe table.
    Ragged rows are padded to the widest row. Returns "" for anything that
    doesn't amount to a real table (no rows, or an all-blank header)."""
    if not rows:
        return ""
    cleaned = [[_clean_cell(c) for c in row] for row in rows if row]
    if not cleaned:
        return ""
    ncols = max(len(r) for r in cleaned)
    if ncols < 1:
        return ""
    cleaned = [r + [""] * (ncols - len(r)) for r in cleaned]
    if not any(cleaned[0]):
        return ""  # a header row with no text in any cell isn't a real table
    lines = ["| " + " | ".join(_escape_pipe(c) for c in cleaned[0]) + " |",
             "| " + " | ".join(["---"] * ncols) + " |"]
    for row in cleaned[1:]:
        lines.append("| " + " | ".join(_escape_pipe(c) for c in row) + " |")
    return "\n".join(lines)


def _render_pdf_page(page, page_number: int) -> str:
    """Interleave a PDF page's prose text and tables in reading order,
    rendering each detected table as GFM. Tables with fewer than 2 rows are
    treated as text, not structure (filters pdfplumber false-positives from
    incidental character spacing, e.g. a parenthetical on the title page —
    a real rule table always has a header + at least one data row).

    A ``<!-- pdf:page=N -->`` marker precedes the page's content so kb.py
    can recover source-page traceability (KB-07) from the stored Markdown
    itself, without a second out-of-band index."""
    marker = f"<!-- pdf:page={page_number} -->"
    try:
        tables = [t for t in page.find_tables() if len(t.rows) >= 2]
    except Exception:  # noqa: BLE001 — table detection is best-effort
        tables = []
    if not tables:
        return f"{marker}\n{page.extract_text() or ''}"

    tables_sorted = sorted(tables, key=lambda t: t.bbox[1])
    segments: list[str] = []
    cursor_top = 0.0
    page_bottom = page.height
    for t in tables_sorted:
        _, top, _, bottom = t.bbox
        top = max(top, cursor_top)
        if top > cursor_top + 1:
            band_text = (page.crop((0, cursor_top, page.width, top)).extract_text() or "").strip()
            if band_text:
                segments.append(band_text)
        try:
            rows = t.extract()
        except Exception:  # noqa: BLE001
            rows = None
        gfm = _rows_to_gfm(rows) if rows else ""
        if gfm:
            segments.append(gfm)
        cursor_top = max(cursor_top, bottom)
    if cursor_top < page_bottom - 1:
        band_text = (page.crop((0, cursor_top, page.width, page_bottom)).extract_text() or "").strip()
        if band_text:
            segments.append(band_text)
    return marker + "\n" + "\n\n".join(segments)


def _extract_docx(content: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ConversionError("python-docx is not installed — cannot read .docx") from exc
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 — any parse failure is a corrupt/invalid file
        raise ConversionError(f"Could not read this .docx file — it may be corrupt: {exc}") from exc
    lines = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading 1"):
            lines.append(f"# {text}")
        elif style.startswith("heading 2"):
            lines.append(f"## {text}")
        elif style.startswith("heading 3"):
            lines.append(f"### {text}")
        else:
            lines.append(text)
    for table in document.tables:
        rows = [[c.text for c in row.cells] for row in table.rows]
        gfm = _rows_to_gfm(rows)
        if gfm:
            lines.append(gfm)
    return "\n\n".join(lines)


def _extract_pdf(content: bytes) -> tuple[str, list[str]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ConversionError("pdfplumber is not installed — cannot read .pdf") from exc
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            # Plain per-page text drives the scanned/image-only density check
            # (unchanged invariant: a page's characters are on the page
            # whether or not they sit inside a detected table, so this check
            # is unaffected by table-aware rendering below). The rendered
            # pages are the table-aware Markdown that actually gets stored.
            plain_pages = [page.extract_text() or "" for page in pdf.pages]
            rendered_pages = [_render_pdf_page(page, i + 1) for i, page in enumerate(pdf.pages)]
    except ConversionError:
        raise
    except Exception as exc:  # noqa: BLE001 — any parse failure is a corrupt/invalid file
        raise ConversionError(f"Could not read this .pdf file — it may be corrupt: {exc}") from exc
    if not plain_pages:
        raise ConversionError("The PDF has no pages.")
    total_chars = sum(len(p) for p in plain_pages)
    if total_chars < _PDF_MIN_CHARS_PER_PAGE * len(plain_pages):
        raise ConversionError(
            "This PDF appears to be scanned/image-only (no extractable text layer). "
            "OCR is out of scope — upload a text-based PDF, DOCX, or Markdown/TXT instead.")
    warnings = []
    empty_pages = [i + 1 for i, p in enumerate(plain_pages) if len((p or "").strip()) < 5]
    if empty_pages:
        warnings.append(f"Page(s) with little or no extractable text: {empty_pages}")
    return "\n\n".join(rendered_pages), warnings


def convert_to_markdown(filename: str, media_type: str, content: bytes) -> dict:
    """Returns {markdown, markdown_sha256, converter_name, converter_version,
    warnings}. Raises ConversionError for unsupported/corrupt/image-only input
    — the caller must not persist a document version on this exception."""
    if len(content) > MAX_BYTES:
        raise ConversionError(f"File exceeds the {MAX_BYTES // (1024 * 1024)} MB limit.")
    if not content:
        raise ConversionError("The uploaded file is empty.")
    resolved_type = sniff_media_type(filename, media_type)
    warnings: list[str] = []
    if resolved_type == "text/plain":
        try:
            markdown = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ConversionError("The .txt file is not valid UTF-8 text.") from exc
        converter = "passthrough-txt"
    elif resolved_type == "text/markdown":
        try:
            markdown = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ConversionError("The .md file is not valid UTF-8 text.") from exc
        converter = "passthrough-md"
    elif resolved_type.endswith("wordprocessingml.document"):
        markdown = _extract_docx(content)
        converter = "python-docx"
    elif resolved_type == "application/pdf":
        markdown, pdf_warnings = _extract_pdf(content)
        warnings.extend(pdf_warnings)
        converter = "pdfplumber"
    else:
        raise ConversionError(f"Unsupported media type: {resolved_type!r}")
    if not markdown.strip():
        raise ConversionError("No readable content was extracted from this file.")
    return {
        "markdown": markdown,
        "markdown_sha256": hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
        "converter_name": converter,
        "converter_version": CONVERTER_VERSION,
        "warnings": warnings,
    }
